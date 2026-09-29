"""Office document readers: DOCX (python-docx), PPTX (python-pptx), XLSX (openpyxl)."""

from __future__ import annotations

from pathlib import Path
from typing import Any, ClassVar

import docx
import openpyxl
import pptx

from embedcraft.domain.entities import CanonicalDocument
from embedcraft.domain.exceptions import DocumentError


class DocxReader:
    SUPPORTED_EXTENSIONS: ClassVar[set[str]] = {".docx"}

    def supports(self, extension: str, mime_type: str = "") -> bool:
        ext = extension.lower()
        if not ext.startswith("."):
            ext = f".{ext}"
        return ext in self.SUPPORTED_EXTENSIONS or "wordprocessingml.document" in mime_type

    def extract(
        self,
        file_path: Path,
        document_id: str,
        metadata: dict[str, Any] | None = None,
    ) -> CanonicalDocument:
        try:
            doc = docx.Document(file_path)
        except Exception as e:
            raise DocumentError(
                message=f"No se pudo abrir el documento DOCX: {file_path.name}",
                technical_detail=str(e),
            ) from e

        paragraphs: list[str] = []
        sections: list[dict[str, Any]] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)
                if p.style and p.style.name.startswith("Heading"):
                    sections.append({"title": text, "level": p.style.name})

        # Extract tables
        tables_data: list[dict[str, Any]] = []
        for t_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                if any(row_cells):
                    table_rows.append(row_cells)
            if table_rows:
                tables_data.append({"table_index": t_idx + 1, "rows": table_rows})
                # Add table representation to text
                table_str = "\n".join([" | ".join(r) for r in table_rows])
                paragraphs.append(f"\n[Tabla {t_idx + 1}]\n{table_str}\n")

        return CanonicalDocument(
            document_id=document_id,
            title=file_path.stem,
            text="\n\n".join(paragraphs),
            sections=sections,
            tables=tables_data,
            metadata=metadata or {},
            source_locator=str(file_path),
        )


class PptxReader:
    SUPPORTED_EXTENSIONS: ClassVar[set[str]] = {".pptx"}

    def supports(self, extension: str, mime_type: str = "") -> bool:
        ext = extension.lower()
        if not ext.startswith("."):
            ext = f".{ext}"
        return ext in self.SUPPORTED_EXTENSIONS or "presentationml.presentation" in mime_type

    def extract(
        self,
        file_path: Path,
        document_id: str,
        metadata: dict[str, Any] | None = None,
    ) -> CanonicalDocument:
        try:
            prs = pptx.Presentation(file_path)
        except Exception as e:
            raise DocumentError(
                message=f"No se pudo abrir la presentación PPTX: {file_path.name}",
                technical_detail=str(e),
            ) from e

        slides_text: list[str] = []
        sections: list[dict[str, Any]] = []

        for idx, slide in enumerate(prs.slides):
            slide_parts = []
            title = f"Diapositiva {idx + 1}"
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        txt = paragraph.text.strip()
                        if txt:
                            slide_parts.append(txt)
            if slide_parts:
                slides_text.append(f"--- Diapositiva {idx + 1} ---\n" + "\n".join(slide_parts))
                sections.append({"title": title, "page": idx + 1})

        return CanonicalDocument(
            document_id=document_id,
            title=file_path.stem,
            text="\n\n".join(slides_text),
            sections=sections,
            tables=[],
            metadata={**(metadata or {}), "slide_count": len(prs.slides)},
            source_locator=str(file_path),
        )


class XlsxReader:
    SUPPORTED_EXTENSIONS: ClassVar[set[str]] = {".xlsx"}

    def supports(self, extension: str, mime_type: str = "") -> bool:
        ext = extension.lower()
        if not ext.startswith("."):
            ext = f".{ext}"
        return ext in self.SUPPORTED_EXTENSIONS or "spreadsheetml.sheet" in mime_type

    def extract(
        self,
        file_path: Path,
        document_id: str,
        metadata: dict[str, Any] | None = None,
    ) -> CanonicalDocument:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        except Exception as e:
            raise DocumentError(
                message=f"No se pudo abrir la hoja de cálculo XLSX: {file_path.name}",
                technical_detail=str(e),
            ) from e

        sheets_text: list[str] = []
        sections: list[dict[str, Any]] = []

        try:
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                rows_text = []
                for row in ws.iter_rows(values_only=True):
                    # Filter out empty rows
                    cells = [str(c).strip() if c is not None else "" for c in row]
                    if any(cells):
                        rows_text.append(" | ".join(cells))
                if rows_text:
                    sheets_text.append(f"--- Hoja: {sheetname} ---\n" + "\n".join(rows_text))
                    sections.append({"title": f"Hoja: {sheetname}"})
        finally:
            wb.close()

        return CanonicalDocument(
            document_id=document_id,
            title=file_path.stem,
            text="\n\n".join(sheets_text),
            sections=sections,
            tables=[],
            metadata={**(metadata or {}), "sheet_names": wb.sheetnames},
            source_locator=str(file_path),
        )
