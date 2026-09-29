"""Unit tests for all supported document format readers."""

import json
from pathlib import Path

import docx
import openpyxl
import pptx
import pymupdf

from embedcraft.adapters.documents.html_reader import HTMLReader
from embedcraft.adapters.documents.office_reader import DocxReader, PptxReader, XlsxReader
from embedcraft.adapters.documents.pdf_reader import PDFReader
from embedcraft.adapters.documents.registry import default_reader_registry
from embedcraft.adapters.documents.structured_reader import CSVReader, JSONReader
from embedcraft.adapters.documents.text_reader import TextReader


def test_text_and_markdown_reader(tmp_path: Path):
    reader = TextReader()
    md_file = tmp_path / "guide.md"
    md_file.write_text("# Guía de Usuario\nBienvenido a EmbedCraft.\n## Sección 1\nDetalles.", encoding="utf-8")

    assert reader.supports(".md")
    assert reader.supports(".txt")

    doc = reader.extract(md_file, document_id="doc-txt")
    assert doc.title == "guide"
    assert "Bienvenido a EmbedCraft." in doc.text
    assert len(doc.sections) == 2
    assert doc.sections[0]["title"] == "Guía de Usuario"


def test_pdf_reader(tmp_path: Path):
    # Create a small valid PDF using PyMuPDF
    pdf_path = tmp_path / "test.pdf"
    doc_fitz = pymupdf.open()
    page = doc_fitz.new_page()
    page.insert_text((50, 50), "Este es un documento PDF de prueba para EmbedCraft.")
    doc_fitz.save(str(pdf_path))
    doc_fitz.close()

    reader = PDFReader()
    assert reader.supports(".pdf")
    doc = reader.extract(pdf_path, document_id="doc-pdf")
    assert "Página 1" in doc.text
    assert "EmbedCraft" in doc.text
    assert doc.metadata["page_count"] == 1


def test_docx_reader(tmp_path: Path):
    docx_path = tmp_path / "test.docx"
    doc_obj = docx.Document()
    doc_obj.add_heading("Título Documento", level=1)
    doc_obj.add_paragraph("Párrafo de prueba en DOCX.")
    table = doc_obj.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Col A"
    table.cell(0, 1).text = "Col B"
    table.cell(1, 0).text = "Val 1"
    table.cell(1, 1).text = "Val 2"
    doc_obj.save(str(docx_path))

    reader = DocxReader()
    assert reader.supports(".docx")
    doc = reader.extract(docx_path, document_id="doc-docx")
    assert "Título Documento" in doc.text
    assert "Párrafo de prueba en DOCX." in doc.text
    assert "Col A | Col B" in doc.text
    assert len(doc.tables) == 1


def test_pptx_reader(tmp_path: Path):
    pptx_path = tmp_path / "test.pptx"
    prs = pptx.Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Presentación de Prueba"
    slide.placeholders[1].text = "Subtítulo explicativo"
    prs.save(str(pptx_path))

    reader = PptxReader()
    assert reader.supports(".pptx")
    doc = reader.extract(pptx_path, document_id="doc-pptx")
    assert "Presentación de Prueba" in doc.text
    assert "Subtítulo explicativo" in doc.text


def test_xlsx_reader(tmp_path: Path):
    xlsx_path = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventas"
    ws.append(["Producto", "Cantidad", "Precio"])
    ws.append(["Laptop", 5, 1200])
    wb.save(str(xlsx_path))
    wb.close()

    reader = XlsxReader()
    assert reader.supports(".xlsx")
    doc = reader.extract(xlsx_path, document_id="doc-xlsx")
    assert "Hoja: Ventas" in doc.text
    assert "Laptop | 5 | 1200" in doc.text


def test_csv_reader(tmp_path: Path):
    csv_path = tmp_path / "data.csv"
    csv_path.write_text("id,nombre,rol\n1,Carlos,Admin\n2,Ana,Dev\n", encoding="utf-8")

    reader = CSVReader()
    assert reader.supports(".csv")
    doc = reader.extract(csv_path, document_id="doc-csv")
    assert "Columnas: id | nombre | rol" in doc.text
    assert "Carlos" in doc.text


def test_json_and_jsonl_reader(tmp_path: Path):
    reader = JSONReader()
    assert reader.supports(".json")
    assert reader.supports(".jsonl")

    json_path = tmp_path / "data.json"
    json_path.write_text(json.dumps([{"item": "A", "val": 10}, {"item": "B", "val": 20}]), encoding="utf-8")
    doc_json = reader.extract(json_path, document_id="doc-json")
    assert "Elemento 1" in doc_json.text
    assert '"item": "A"' in doc_json.text

    jsonl_path = tmp_path / "data.jsonl"
    jsonl_path.write_text('{"user": "u1"}\n{"user": "u2"}\n', encoding="utf-8")
    doc_jsonl = reader.extract(jsonl_path, document_id="doc-jsonl")
    assert "Registro 1" in doc_jsonl.text
    assert '"user": "u1"' in doc_jsonl.text


def test_html_reader(tmp_path: Path):
    html_path = tmp_path / "page.html"
    html_path.write_text(
        "<html><head><title>Título HTML</title><script>alert('hack');</script></head>"
        "<body><h1>Encabezado Principal</h1><p>Contenido del párrafo web.</p></body></html>",
        encoding="utf-8",
    )

    reader = HTMLReader()
    assert reader.supports(".html")
    doc = reader.extract(html_path, document_id="doc-html")
    assert doc.title == "Título HTML"
    assert "Encabezado Principal" in doc.text
    assert "Contenido del párrafo web." in doc.text
    assert "alert('hack')" not in doc.text  # Scripts safely stripped


def test_registry_detection(tmp_path: Path):
    txt_file = tmp_path / "a.txt"
    txt_file.touch()
    r = default_reader_registry.get_reader_for_file(txt_file)
    assert isinstance(r, TextReader)
